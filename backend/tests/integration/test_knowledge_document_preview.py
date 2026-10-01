"""
Knowledge-base documents are previewed and downloaded in their own format.

Uploads used to keep only the extracted text, so "Download" on a PDF handed
back ``guide.pdf.txt``. The upload is now kept byte-for-byte in
``document_files``; download returns it unchanged, and preview says how to
render it. Pasted text and pre-existing documents still work from the text.
"""
import io
import uuid
from datetime import datetime, timedelta
from decimal import Decimal

import httpx
import pytest
import pytest_asyncio
from httpx import ASGITransport
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker

from app.main import app
from app.database import get_db
from fastapi import Depends

from app.core.dependencies import get_current_user
from app.models.user import User, Organization, OrganizationMember
from app.models.knowledge_base import KnowledgeBase, Document, DocumentFile
from app.models.subscription import Subscription, SubscriptionPlan
from app.services.billing import catalog
from app.services.knowledge_base.rag_service import RAGService

pytestmark = pytest.mark.asyncio


def _pdf_bytes(text: str) -> bytes:
    """A minimal one-page PDF with extractable text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Service", "Price"])
    ws.append(["Cleaning", 80])
    ws.append(["Whitening", 250.5])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest_asyncio.fixture
async def ctx(db_engine, db_session, monkeypatch):
    # Embedding runs in the background against OpenAI; it is not under test.
    async def _no_processing(self, *_a, **_k):
        return None

    monkeypatch.setattr(RAGService, "_process_document_bg", _no_processing)

    async def _owner(tag: str) -> tuple:
        user = User(email=f"{tag}-{uuid.uuid4().hex[:6]}@example.com", hashed_password="x", full_name=tag, is_active=True)
        db_session.add(user)
        await db_session.flush()
        org = Organization(name=tag, slug=f"{tag}-{uuid.uuid4().hex[:8]}", owner_id=user.id, is_active=True)
        db_session.add(org)
        await db_session.flush()
        db_session.add(OrganizationMember(organization_id=org.id, user_id=user.id, role="owner"))
        return user, org

    user, org = await _owner("pearl")
    plan = SubscriptionPlan(
        name="Growth", slug=f"growth-{uuid.uuid4().hex[:6]}",
        stripe_product_id="prod_test", stripe_price_id=f"price_{uuid.uuid4().hex[:6]}",
        price_monthly=Decimal("99.00"), entitlements=catalog.entitlements_for_plan("growth"),
    )
    db_session.add(plan)
    await db_session.flush()
    now = datetime.utcnow()
    db_session.add(Subscription(
        organization_id=org.id, plan_id=plan.id, status="active", billing_period="monthly",
        current_period_start=now, current_period_end=now + timedelta(days=30),
    ))
    outsider, _ = await _owner("other")
    kb = KnowledgeBase(organization_id=org.id, name="Clinic")
    db_session.add(kb)
    await db_session.commit()

    sessionmaker = async_sessionmaker(db_engine, expire_on_commit=False)

    async def _db():
        async with sessionmaker() as session:
            yield session

    acting = {"id": user.id}

    async def _current_user(db=Depends(get_db)):
        return (await db.execute(select(User).where(User.id == acting["id"]))).scalar_one()

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = _current_user
    async with httpx.AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield {"client": client, "kb_id": kb.id, "acting": acting, "outsider": outsider.id, "db": db_session}
    app.dependency_overrides.clear()


async def _upload(ctx, filename: str, data: bytes, content_type: str = "application/octet-stream") -> str:
    res = await ctx["client"].post(
        "/api/v1/knowledge/documents/upload",
        data={"knowledge_base_id": str(ctx["kb_id"])},
        files={"file": (filename, data, content_type)},
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


async def test_pdf_downloads_byte_for_byte_and_previews_as_pdf(ctx):
    original = _pdf_bytes("Cleaning costs 80 dollars")
    doc_id = await _upload(ctx, "Price list.pdf", original)

    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/preview")).json()
    assert preview["kind"] == "pdf"
    assert preview["original_available"] is True
    assert preview["text"] is None  # the viewer renders the bytes, not text

    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/download")
    assert res.status_code == 200
    assert res.content == original
    assert res.headers["content-type"] == "application/pdf"
    assert 'filename="Price list.pdf"' in res.headers["content-disposition"]


async def test_xlsx_previews_as_sheets_and_downloads_unchanged(ctx):
    original = _xlsx_bytes()
    doc_id = await _upload(ctx, "prices.xlsx", original)

    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/preview")).json()
    assert preview["kind"] == "sheet"
    assert preview["sheets"][0]["name"] == "Prices"
    assert preview["sheets"][0]["rows"] == [["Service", "Price"], ["Cleaning", "80"], ["Whitening", "250.5"]]

    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/download")
    assert res.content == original
    assert "spreadsheetml" in res.headers["content-type"]


async def test_csv_parses_quoted_cells(ctx):
    doc_id = await _upload(ctx, "hours.csv", b'Day,Hours\nMonday,"9am, to 5pm"\n')
    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/preview")).json()
    assert preview["kind"] == "sheet"
    assert preview["sheets"][0]["rows"] == [["Day", "Hours"], ["Monday", "9am, to 5pm"]]


async def test_markdown_and_json_keep_their_kind(ctx):
    md_id = await _upload(ctx, "guide.md", b"# Pearl Dental\n\n- Open 9-5\n")
    js_id = await _upload(ctx, "faq.json", b'{"q": "Parking?", "a": "Free"}')

    md = (await ctx["client"].get(f"/api/v1/knowledge/documents/{md_id}/preview")).json()
    js = (await ctx["client"].get(f"/api/v1/knowledge/documents/{js_id}/preview")).json()
    assert (md["kind"], md["text"]) == ("markdown", "# Pearl Dental\n\n- Open 9-5\n")
    assert js["kind"] == "json"

    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{md_id}/download")
    assert res.content == b"# Pearl Dental\n\n- Open 9-5\n"
    assert res.headers["content-type"].startswith("text/markdown")


async def test_document_without_original_falls_back_to_extracted_text(ctx):
    """Pasted text, and PDFs uploaded before originals were kept."""
    doc = Document(
        knowledge_base_id=ctx["kb_id"], title="old-guide.pdf", content="Extracted words",
        content_hash=uuid.uuid4().hex, source_type="file", processing_status="completed",
    )
    ctx["db"].add(doc)
    await ctx["db"].commit()

    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc.id}/preview")).json()
    assert (preview["kind"], preview["original_available"], preview["text"]) == ("text", False, "Extracted words")

    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc.id}/download")
    assert res.content == b"Extracted words"
    assert 'filename="old-guide.pdf.txt"' in res.headers["content-disposition"]


async def test_older_text_upload_downloads_under_its_own_name(ctx):
    """A .md uploaded before originals were kept: its stored text is the file."""
    doc = Document(
        knowledge_base_id=ctx["kb_id"], title="clinic-guide.md", content="# Guide\n",
        content_hash=uuid.uuid4().hex, source_type="file", processing_status="completed",
    )
    ctx["db"].add(doc)
    await ctx["db"].commit()

    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc.id}/preview")).json()
    assert (preview["kind"], preview["original_available"]) == ("markdown", True)

    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc.id}/download")
    assert res.content == b"# Guide\n"
    assert 'filename="clinic-guide.md"' in res.headers["content-disposition"]
    assert res.headers["content-type"].startswith("text/markdown")


async def test_pasted_text_still_downloads_as_txt(ctx):
    res = await ctx["client"].post("/api/v1/knowledge/documents", json={
        "knowledge_base_id": str(ctx["kb_id"]), "title": "Refund policy", "content": "No refunds after 30 days.",
    })
    assert res.status_code == 201, res.text
    doc_id = res.json()["id"]

    preview = (await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/preview")).json()
    assert (preview["kind"], preview["original_available"]) == ("text", False)
    res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/download")
    assert 'filename="Refund policy.txt"' in res.headers["content-disposition"]


async def test_other_workspace_cannot_preview_or_download(ctx):
    doc_id = await _upload(ctx, "guide.md", b"secret")
    ctx["acting"]["id"] = ctx["outsider"]
    for path in ("preview", "download"):
        res = await ctx["client"].get(f"/api/v1/knowledge/documents/{doc_id}/{path}")
        assert res.status_code == 403


async def test_deleting_the_document_removes_the_original(ctx):
    doc_id = await _upload(ctx, "guide.md", b"bye")
    res = await ctx["client"].delete(f"/api/v1/knowledge/documents/{doc_id}")
    assert res.status_code == 204
    left = (await ctx["db"].execute(
        select(DocumentFile).where(DocumentFile.document_id == uuid.UUID(doc_id))
    )).scalar_one_or_none()
    assert left is None
